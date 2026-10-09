import io
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_docx(report_data: dict) -> io.BytesIO:
    """
    Genera un documento Word (.docx) con la estructura oficial idéntica
    al ejemplo Ayuda_Memoria_Exportaciones_Uruguay.docx
    """
    doc = docx.Document()

    # Configurar márgenes de página (1 pulgada)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    country = report_data.get('country', '').capitalize()
    paragraphs = report_data.get('paragraphs', {})

    # Título Principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run(report_data.get('title', 'Ayuda memoria'))
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(20)
    # Azul institucional oscuro como en el original
    r_title.font.color.rgb = RGBColor(0, 51, 102)

    # Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(report_data.get('subtitle', f"Empresas exportadoras a {country} - Ruta Exportadora"))
    r_sub.bold = True
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(80, 80, 80)

    # Párrafo 1: Estadísticas Generales
    p1 = doc.add_paragraph()
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(8)
    r_p1 = p1.add_run(paragraphs.get('p1', ''))
    r_p1.font.name = 'Calibri'
    r_p1.font.size = Pt(11)

    # Párrafo 2: Productos y Regiones
    p2 = doc.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_after = Pt(12)
    r_p2 = p2.add_run(paragraphs.get('p2', ''))
    r_p2.font.name = 'Calibri'
    r_p2.font.size = Pt(11)

    # Párrafo Encabezado de Eventos
    p_ev_head = doc.add_paragraph()
    p_ev_head.paragraph_format.space_before = Pt(6)
    p_ev_head.paragraph_format.space_after = Pt(6)
    r_ev_head = p_ev_head.add_run(paragraphs.get('events_heading', 'Los principales programas en donde participaron las empresas fueron:'))
    r_ev_head.bold = True
    r_ev_head.font.name = 'Calibri'
    r_ev_head.font.size = Pt(11)

    # Viñetas Top 15 Eventos
    events_bullets = paragraphs.get('events_bullets', [])
    if events_bullets:
        for item_text in events_bullets:
            p_item = doc.add_paragraph(style='List Bullet')
            p_item.paragraph_format.space_before = Pt(0)
            p_item.paragraph_format.space_after = Pt(3)
            p_item.paragraph_format.line_spacing = 1.15
            r_item = p_item.add_run(item_text)
            r_item.font.name = 'Calibri'
            r_item.font.size = Pt(11)
    else:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(3)
        p_item.paragraph_format.line_spacing = 1.15
        r_item = p_item.add_run("No se registraron eventos o programas para las empresas exportadoras a este destino.")
        r_item.font.name = 'Calibri'
        r_item.font.size = Pt(11)

    # Guardar en buffer BytesIO
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
